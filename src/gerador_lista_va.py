import os
import sys
import glob

# Garante que a raiz do projeto esteja no sys.path para qualquer forma de execução
diretorio_raiz = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if diretorio_raiz not in sys.path:
    sys.path.insert(0, diretorio_raiz)

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

try:
    from src.mapeamento import (
        normalizar_texto,
        padronizar_empresa_va,
        carregar_mapa_secretarias,
        identificar_secretaria,
        obter_centro_custo_usecred,
        verificar_troca_centro_custo_usecred
    )
    from src.leitor_pdf import extrair_funcionarios_pdf
    from src.leitor_transferencias import extrair_transferencias_pdf
except ImportError:
    from mapeamento import (
        normalizar_texto,
        padronizar_empresa_va,
        carregar_mapa_secretarias,
        identificar_secretaria,
        obter_centro_custo_usecred,
        verificar_troca_centro_custo_usecred
    )
    from leitor_pdf import extrair_funcionarios_pdf
    from leitor_transferencias import extrair_transferencias_pdf


def carregar_base_inicial(caminho_inicial="data/entrega_inicial"):
    """
    Carrega CPFs, secretarias e operadoras da entrega inicial de abril
    para mapear histórico de quem já existia.
    """
    if not os.path.exists(caminho_inicial):
        return {}
    
    mapa_codigos, mapa_termos = carregar_mapa_secretarias("data/siglas secretarias.ods")
    mapa_servidores = {}
    arquivos = [f for f in sorted(os.listdir(caminho_inicial)) if f.endswith(('.ods', '.xls', '.xlsx'))]
    
    for fname in arquivos:
        fpath = os.path.join(caminho_inicial, fname)
        ext = os.path.splitext(fname)[1].lower()
        engine = "odf" if ext == ".ods" else ("openpyxl" if ext == ".xlsx" else "xlrd")
        
        # Identifica operadora pelo nome do arquivo
        operadora_arquivo = padronizar_empresa_va(fname)
        
        try:
            xl = pd.ExcelFile(fpath, engine=engine)
            for s in xl.sheet_names:
                df = xl.parse(s, header=None)
                header_idx = None
                for idx in range(min(15, len(df))):
                    row_vals = [str(x).strip().upper() for x in df.iloc[idx].values if pd.notna(x)]
                    if "CPF" in row_vals:
                        header_idx = idx
                        break
                
                if header_idx is not None and header_idx + 1 < len(df):
                    cpf_col = None
                    nome_col = None
                    for c in range(df.shape[1]):
                        val = str(df.iloc[header_idx, c]).strip().upper()
                        if val == "CPF":
                            cpf_col = c
                        elif val in ["NOME", "NOME DO FUNCIONÁRIO", "FUNCIONÁRIO"]:
                            nome_col = c
                    
                    # Identifica secretaria da aba
                    sec_aba = "NÃO IDENTIFICADA"
                    if "Educa" in fname:
                        sec_aba = mapa_codigos.get(107, "107 - SECRETARIA MUNICIPAL DE EDUCAÇÃO")
                        operadora_real = padronizar_empresa_va(s)
                    elif "Sa" in fname:
                        sec_aba = mapa_codigos.get(114, "114 - SECRETARIA MUNICIPAL DE SAÚDE")
                        operadora_real = padronizar_empresa_va(s)
                    else:
                        sec_aba = identificar_secretaria(s, mapa_codigos, mapa_termos)
                        operadora_real = operadora_arquivo
                    
                    if cpf_col is not None and nome_col is not None:
                        dados = df.iloc[header_idx + 1:].copy()
                        for _, r in dados.iterrows():
                            cpf = str(r[cpf_col]).strip() if pd.notna(r[cpf_col]) else ""
                            nome = str(r[nome_col]).strip() if pd.notna(r[nome_col]) else ""
                            if nome and nome.upper() != "NOME":
                                n_norm = normalizar_texto(nome)
                                if n_norm not in mapa_servidores:
                                    mapa_servidores[n_norm] = {
                                        "cpf": cpf,
                                        "secretaria_anterior": sec_aba,
                                        "empresa_anterior": operadora_real,
                                        "origem": f"{fname} ({s})"
                                    }
        except Exception:
            pass
            
    return mapa_servidores


def processar_mes(caminho_pdf, mes_ano_referencia="Setembro/2026", output_path=None, caminho_transferencias=None):
    """
    Processa os admitidos do PDF e cruza com a planilha de escolhas e transferências.
    """
    mapa_codigos, mapa_termos = carregar_mapa_secretarias("data/siglas secretarias.ods")
    base_historico = carregar_base_inicial("data/entrega_inicial")
    
    # 1. Extrai do PDF
    servidores_pdf = extrair_funcionarios_pdf(caminho_pdf)
    print(f"-> Servidores encontrados no PDF de admissões: {len(servidores_pdf)}")
    
    # 2. Carrega planilha de escolhas
    caminho_escolhas = os.path.join("data", "novos_e_alteracoes", "Escolhas das empresas de Vale Alimentação.ods")
    if not os.path.exists(caminho_escolhas):
        raise FileNotFoundError(f"Arquivo de escolhas não encontrado: {caminho_escolhas}")
        
    df_escolhas = pd.read_excel(caminho_escolhas, engine="odf")
    validos = df_escolhas[df_escolhas['Nome'].notna() & (df_escolhas['Nome'] != 'Nome')].copy()
    validos['nome_norm'] = validos['Nome'].apply(normalizar_texto)
    
    dict_escolhas = {}
    for _, r in validos.iterrows():
        nn = r['nome_norm']
        dict_escolhas[nn] = {
            "nome_original": r['Nome'],
            "cargo": r['Cargo'],
            "secretaria_raw": r.get('SECRETARIA', None),
            "inicio": r.get('Início', None),
            "empresa_raw": r.get('Empresa do V.A:', None)
        }
        
    # 3. Cruzamento
    lista_final = []
    alertas = []
    
    for s in servidores_pdf:
        nn = normalizar_texto(s['nome'])
        if nn in dict_escolhas:
            escolha = dict_escolhas[nn]
            empresa_padrao = padronizar_empresa_va(escolha['empresa_raw'])
            sec_padrao = identificar_secretaria(escolha['secretaria_raw'], mapa_codigos, mapa_termos)
            
            if empresa_padrao == "Não optou por nenhuma":
                alertas.append({
                    "funcional": s['funcional'],
                    "nome": s['nome'],
                    "secretaria": sec_padrao,
                    "empresa": empresa_padrao,
                    "situacao": "Não optou por nenhuma operadora de Vale Alimentação"
                })
            else:
                # Determina motivo e aplica regra especial da USECRED
                motivo = "Novo Servidor"
                ja_existia = nn in base_historico
                gerar_novo_cartao = True
                
                if empresa_padrao == "USECRED":
                    cc_novo = obter_centro_custo_usecred(sec_padrao)
                    if ja_existia:
                        hist = base_historico[nn]
                        sec_ant = hist.get("secretaria_anterior", "NÃO IDENTIFICADA")
                        emp_ant = hist.get("empresa_anterior", "OUTRA")
                        
                        if emp_ant == "USECRED":
                            trocou_cc, cc_ant, _ = verificar_troca_centro_custo_usecred(sec_ant, sec_padrao)
                            if trocou_cc:
                                motivo = f"USECRED: Troca de Centro de Custo ({cc_ant} -> {cc_novo})"
                            else:
                                # Regra: Mesmo centro de custo na USECRED NÃO gera novo cartão
                                gerar_novo_cartao = False
                                alertas.append({
                                    "funcional": s['funcional'],
                                    "nome": s['nome'],
                                    "secretaria": sec_padrao,
                                    "empresa": empresa_padrao,
                                    "situacao": f"USECRED: Mantém cartão existente. Mesmo Centro de Custo ({cc_novo}) - Sem novo cartão"
                                })
                        else:
                            motivo = f"USECRED: Troca de Operadora ({emp_ant} -> USECRED / {cc_novo})"
                    else:
                        motivo = f"USECRED: Admissão (Centro: {cc_novo})"
                else:
                    if ja_existia:
                        hist = base_historico[nn]
                        emp_ant = hist.get("empresa_anterior", "OUTRA")
                        
                        if emp_ant == empresa_padrao:
                            # Regra: Já possui cartão ativo da mesma operadora, mantém o existente
                            gerar_novo_cartao = False
                            alertas.append({
                                "funcional": s['funcional'],
                                "nome": s['nome'],
                                "secretaria": sec_padrao,
                                "empresa": empresa_padrao,
                                "situacao": f"Mantém cartão existente. Já possui cartão ativo da operadora ({empresa_padrao}) - Sem novo cartão"
                            })
                        else:
                            motivo = f"Troca de Operadora ({emp_ant} -> {empresa_padrao})"
                    else:
                        motivo = "Novo Servidor (Admissão)"

                if gerar_novo_cartao:
                    lista_final.append({
                        "funcional": s['funcional'],
                        "nome": s['nome'],
                        "cargo": s['cargo'],
                        "admissao": s['admissao'],
                        "secretaria": sec_padrao,
                        "empresa_va": empresa_padrao,
                        "historico": motivo
                    })

        else:
            # Servidor no PDF mas que não consta na planilha de escolhas
            alertas.append({
                "funcional": s['funcional'],
                "nome": s['nome'],
                "secretaria": "NÃO LOCALIZADA (Sem escolha cadastrada)",
                "empresa": "PENDENTE",
                "situacao": "Consta no relatório de admissão mas NÃO preencheu escolha de operadora"
            })
            
    # 3.1 Cruzamento com Relatório de Transferências (Regra Especial USECRED)
    if caminho_transferencias:
        if isinstance(caminho_transferencias, str):
            lista_arquivos_transf = [caminho_transferencias]
        else:
            lista_arquivos_transf = list(caminho_transferencias)
    else:
        transf_files = glob.glob(os.path.join("data", "*tranfer*.pdf")) + glob.glob(os.path.join("data", "*transfer*.pdf"))
        lista_arquivos_transf = sorted(list(set(transf_files)))

    transferencias = []
    chaves_transf_vistas = set()
    for arq_transf in lista_arquivos_transf:
        if os.path.exists(arq_transf):
            t_extraidas = extrair_transferencias_pdf(arq_transf)
            for t_item in t_extraidas:
                chave = (t_item['funcional'], t_item['dt_ini'], t_item['sec_ini_cod'])
                if chave not in chaves_transf_vistas:
                    chaves_transf_vistas.add(chave)
                    transferencias.append(t_item)

    if transferencias:
        print(f"-> Total de transferências carregadas (desduplicadas): {len(transferencias)}")
        
        meses_dict = {
            "JANEIRO": 1, "FEVEREIRO": 2, "MARCO": 3, "MARÇO": 3, "ABRIL": 4,
            "MAIO": 5, "JUNHO": 6, "JULHO": 7, "AGOSTO": 8,
            "SETEMBRO": 9, "OUTUBRO": 10, "NOVEMBRO": 11, "DEZEMBRO": 12
        }
        meses_nomes = {
            1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
            5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
            9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro"
        }
        
        ref_mes = 9
        ref_ano = 2026
        partes_ref = mes_ano_referencia.replace("-", "/").split("/")
        if len(partes_ref) == 2:
            m_str = normalizar_texto(partes_ref[0])
            a_str = partes_ref[1].strip()
            ref_mes = int(m_str) if m_str.isdigit() else meses_dict.get(m_str, 9)
            if a_str.isdigit():
                ref_ano = int(a_str)
                
        funcionais_ja_na_lista = {s['funcional'] for s in lista_final}
        funcionais_ja_nos_alertas = {a['funcional'] for a in alertas}
        
        for t in transferencias:
            nn = normalizar_texto(t['nome'])
            hist = base_historico.get(nn)
            if not hist:
                continue
                
            empresa_hist = hist.get("empresa_anterior", "")
            
            # Regra para USECRED em transferências
            if empresa_hist == "USECRED" and t['sec_fim_cod'] != t['sec_ini_cod']:
                sec_orig = mapa_codigos.get(t['sec_fim_cod'], f"Sec {t['sec_fim_cod']}")
                sec_dest = mapa_codigos.get(t['sec_ini_cod'], f"Sec {t['sec_ini_cod']}")
                trocou_cc, cc_ant, cc_novo = verificar_troca_centro_custo_usecred(sec_orig, sec_dest)
                
                if trocou_cc and t['dt_ini_dt']:
                    t_mes = t['dt_ini_dt'].month
                    t_ano = t['dt_ini_dt'].year
                    
                    # Regra do mês seguinte para entrega do cartão USECRED
                    m_entrega = t_mes + 1 if t_mes < 12 else 1
                    a_entrega = t_ano if t_mes < 12 else t_ano + 1
                    nome_mes_entrega = meses_nomes.get(m_entrega, str(m_entrega))
                    
                    # Caso 1: Mês atual de processamento é o mês da transferência
                    # (Cartão ainda não chegou -> entra em Alertas com previsão para o próximo mês)
                    if t_mes == ref_mes and t_ano == ref_ano:
                        if t['funcional'] not in funcionais_ja_nos_alertas:
                            alertas.append({
                                "funcional": t['funcional'],
                                "nome": t['nome'],
                                "secretaria": sec_dest,
                                "empresa": "USECRED",
                                "situacao": f"USECRED: Transferido em {t['dt_ini']} ({cc_ant} -> {cc_novo}). Troca de Centro de Custo programada para entrega em {nome_mes_entrega}/{a_entrega} (mês seguinte à transferência)"
                            })
                            funcionais_ja_nos_alertas.add(t['funcional'])
                            
                    # Caso 2: Mês atual de processamento é exatamente o mês da entrega prevista
                    # (Chegou o mês seguinte -> entra na folha de entrega imediata)
                    elif ref_mes == m_entrega and ref_ano == a_entrega:
                        if t['funcional'] not in funcionais_ja_na_lista:
                            lista_final.append({
                                "funcional": t['funcional'],
                                "nome": t['nome'],
                                "cargo": "TRANSFERÊNCIA",
                                "admissao": t['dt_ini'],
                                "secretaria": sec_dest,
                                "empresa_va": "USECRED",
                                "historico": f"USECRED: Troca de Centro de Custo ({cc_ant} -> {cc_novo}) - Transferência de {t['dt_ini']}"
                            })
                            funcionais_ja_na_lista.add(t['funcional'])

    print(f"-> Servidores elegíveis para novos cartões: {len(lista_final)}")
    print(f"-> Alertas / Pendências detectados: {len(alertas)}")
    
    # 4. Geração do Excel formatado
    if output_path is None:
        os.makedirs("output", exist_ok=True)
        mes_slug = mes_ano_referencia.replace("/", "-").lower()
        output_path = os.path.join("output", f"lista_VA_{mes_slug}.xlsx")
        
    gerar_excel(lista_final, alertas, mes_ano_referencia, output_path)
    print(f"-> Arquivo Excel gerado com sucesso em: {output_path}")
    return output_path

def gerar_excel(lista_final, alertas, mes_ano_ref, caminho_saida):
    wb = openpyxl.Workbook()
    # Remove sheet padrão
    wb.remove(wb.active)
    
    # Estilos reutilizáveis
    font_titulo = Font(name="Tahoma", size=16, bold=True)
    font_subtitulo = Font(name="Tahoma", size=14, bold=True)
    font_texto_pequeno = Font(name="Tahoma", size=10, bold=False)
    font_sec_titulo = Font(name="Calibri", size=11, bold=True)
    font_th = Font(name="Calibri", size=11, bold=True)
    font_td = Font(name="Calibri", size=11, bold=False)
    
    thin_border = Border(
        left=Side(style='thin', color='A0A0A0'),
        right=Side(style='thin', color='A0A0A0'),
        top=Side(style='thin', color='A0A0A0'),
        bottom=Side(style='thin', color='A0A0A0')
    )
    fill_th = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_sec = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    
    # ==========================================
    # ABA 1: Consulta (Modelo de Entrega)
    # ==========================================
    ws_consulta = wb.create_sheet(title="Consulta")
    ws_consulta.views.sheetView[0].showGridLines = True
    
    # Larguras de coluna
    ws_consulta.column_dimensions['A'].width = 13.0
    ws_consulta.column_dimensions['B'].width = 50.0
    ws_consulta.column_dimensions['C'].width = 18.0
    ws_consulta.column_dimensions['D'].width = 44.0
    
    # Cabeçalho Institucional
    ws_consulta.merge_cells('A1:D1')
    ws_consulta['A1'] = "PREFEITURA DO MUNICÍPIO DE PIRACICABA"
    ws_consulta['A1'].font = font_titulo
    ws_consulta['A1'].alignment = align_center
    
    ws_consulta.merge_cells('A2:D2')
    ws_consulta['A2'] = "Departamento de Recursos Humanos"
    ws_consulta['A2'].font = font_subtitulo
    ws_consulta['A2'].alignment = align_center
    
    ws_consulta['A4'] = "Relação de funcionários para entrega do cartão de Vale Alimentação"
    ws_consulta['A4'].font = font_texto_pequeno
    
    ws_consulta['A5'] = f"Referente: {mes_ano_ref}"
    ws_consulta['A5'].font = font_texto_pequeno
    
    # Agrupa servidores por secretaria
    df_final = pd.DataFrame(lista_final)
    linha_atual = 8
    
    if not df_final.empty:
        # Ordena por Secretaria e depois por Nome
        df_final['sec_str'] = df_final['secretaria'].astype(str)
        secretarias_ordenadas = sorted(df_final['sec_str'].unique())
        
        for sec in secretarias_ordenadas:
            sub_df = df_final[df_final['sec_str'] == sec].sort_values(by="nome")
            
            # Título da Secretaria
            ws_consulta.cell(row=linha_atual, column=1, value=sec).font = font_sec_titulo
            linha_atual += 1
            
            # Cabeçalho da Tabela
            headers = ["Funcional", "Nome", "Data Retirada", "empresa VA"]
            for col_idx, h_text in enumerate(headers, start=1):
                c = ws_consulta.cell(row=linha_atual, column=col_idx, value=h_text)
                c.font = font_th
                c.fill = fill_th
                c.border = thin_border
                c.alignment = align_center if col_idx in [1, 3] else align_left
            linha_atual += 1
            
            # Linhas dos Servidores
            for _, row_srv in sub_df.iterrows():
                # Coluna A: Funcional
                cA = ws_consulta.cell(row=linha_atual, column=1, value=row_srv['funcional'])
                cA.font = font_td
                cA.alignment = align_center
                cA.border = thin_border
                
                # Coluna B: Nome
                cB = ws_consulta.cell(row=linha_atual, column=2, value=row_srv['nome'])
                cB.font = font_td
                cB.alignment = align_left
                cB.border = thin_border
                
                # Coluna C: Data Retirada (em branco para visto/assinatura)
                cC = ws_consulta.cell(row=linha_atual, column=3, value=None)
                cC.font = font_td
                cC.alignment = align_center
                cC.border = thin_border
                
                # Coluna D: Empresa VA
                cD = ws_consulta.cell(row=linha_atual, column=4, value=row_srv['empresa_va'])
                cD.font = font_td
                cD.alignment = align_left
                cD.border = thin_border
                
                linha_atual += 1
                
            # Espaço entre secretarias
            linha_atual += 2
    else:
        ws_consulta.cell(row=8, column=1, value="Nenhum servidor com direito a novo cartão neste mês.").font = font_sec_titulo

    # ==========================================
    # ABA 2: Planilha1 (Base de Controle)
    # ==========================================
    ws_plan1 = wb.create_sheet(title="Planilha1")
    ws_plan1.views.sheetView[0].showGridLines = True
    
    headers_plan1 = ["Funcional", "Nome", "secretaria", "empresa VA", "recebido", "Motivo / Histórico"]
    for c_idx, h_text in enumerate(headers_plan1, start=1):
        c = ws_plan1.cell(row=1, column=c_idx, value=h_text)
        c.font = font_th
        c.fill = fill_th
        c.border = thin_border
        c.alignment = align_center if c_idx in [1, 5] else align_left
        
    for r_idx, row_srv in enumerate(lista_final, start=2):
        c1 = ws_plan1.cell(row=r_idx, column=1, value=row_srv['funcional'])
        c1.alignment = align_center
        c1.font = font_td
        c1.border = thin_border
        
        c2 = ws_plan1.cell(row=r_idx, column=2, value=row_srv['nome'])
        c2.alignment = align_left
        c2.font = font_td
        c2.border = thin_border
        
        c3 = ws_plan1.cell(row=r_idx, column=3, value=row_srv['secretaria'])
        c3.alignment = align_left
        c3.font = font_td
        c3.border = thin_border
        
        c4 = ws_plan1.cell(row=r_idx, column=4, value=row_srv['empresa_va'])
        c4.alignment = align_left
        c4.font = font_td
        c4.border = thin_border
        
        c5 = ws_plan1.cell(row=r_idx, column=5, value=None)
        c5.alignment = align_center
        c5.font = font_td
        c5.border = thin_border

        c6 = ws_plan1.cell(row=r_idx, column=6, value=row_srv.get('historico', 'Novo Servidor'))
        c6.alignment = align_left
        c6.font = font_td
        c6.border = thin_border

    # Ajusta larguras da Planilha1
    ws_plan1.column_dimensions['A'].width = 14.0
    ws_plan1.column_dimensions['B'].width = 48.0
    ws_plan1.column_dimensions['C'].width = 45.0
    ws_plan1.column_dimensions['D'].width = 30.0
    ws_plan1.column_dimensions['E'].width = 12.0
    ws_plan1.column_dimensions['F'].width = 45.0


    # ==========================================
    # ABA 3: Pendências e Alertas
    # ==========================================
    ws_alertas = wb.create_sheet(title="Pendências e Alertas")
    ws_alertas.views.sheetView[0].showGridLines = True
    
    headers_alertas = ["Funcional", "Nome", "Secretaria", "Empresa", "Situação / Ação Recomendada"]
    for c_idx, h_text in enumerate(headers_alertas, start=1):
        c = ws_alertas.cell(row=1, column=c_idx, value=h_text)
        c.font = font_th
        c.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Tom amarelo claro para alertas
        c.border = thin_border
        c.alignment = align_center if c_idx == 1 else align_left
        
    for r_idx, alt in enumerate(alertas, start=2):
        c1 = ws_alertas.cell(row=r_idx, column=1, value=alt['funcional'])
        c1.alignment = align_center
        c1.font = font_td
        c1.border = thin_border
        
        c2 = ws_alertas.cell(row=r_idx, column=2, value=alt['nome'])
        c2.font = font_td
        c2.border = thin_border
        
        c3 = ws_alertas.cell(row=r_idx, column=3, value=alt['secretaria'])
        c3.font = font_td
        c3.border = thin_border
        
        c4 = ws_alertas.cell(row=r_idx, column=4, value=alt['empresa'])
        c4.font = font_td
        c4.border = thin_border
        
        c5 = ws_alertas.cell(row=r_idx, column=5, value=alt['situacao'])
        c5.font = font_td
        c5.border = thin_border
        
    ws_alertas.column_dimensions['A'].width = 14.0
    ws_alertas.column_dimensions['B'].width = 45.0
    ws_alertas.column_dimensions['C'].width = 35.0
    ws_alertas.column_dimensions['D'].width = 25.0
    ws_alertas.column_dimensions['E'].width = 65.0

    wb.save(caminho_saida)
    return caminho_saida

if __name__ == "__main__":
    import glob
    pdfs = glob.glob(os.path.join(diretorio_raiz, "data", "*.pdf"))
    if pdfs:
        pdf_selecionado = sorted(pdfs)[-1]
        print(f"Executando diretamente para o PDF: {os.path.basename(pdf_selecionado)}")
        processar_mes(pdf_selecionado, "Setembro/2026")
    else:
        print("Nenhum PDF encontrado em data/")

